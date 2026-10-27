# -*- coding: utf-8 -*-
"""Создаёт демо-шаблон template.pptx — замену «ручной» презентации компании.

В реальном проекте шаблоном становится ваша собственная презентация в корпоративных цветах:
1) числа и тексты, которые меняются каждый месяц, заменяют метками вида {{период}};
2) таблицам и графикам дают имена (в PowerPoint: Главная → Выделить → Область выделения);
3) в ячейках таблиц оставляют «—»: скрипт заменит прочерк, сохранив шрифт и цвет ячейки.
Скрипт сборки ничего не рисует сам — он только заполняет шаблон, поэтому брендбук остаётся в шаблоне.
Запуск: python make_template.py
"""
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent

# «демо-брендбук»: палитра курса вместо корпоративных цветов
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
TEAL = RGBColor(0x2A, 0x9D, 0x8F)
GOLD = RGBColor(0xE9, 0xC4, 0x6A)
GREY = RGBColor(0x6B, 0x72, 0x80)
LIGHT = RGBColor(0xE8, 0xEE, 0xF6)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x2B, 0x2D, 0x33)
DIRECTIONS = ["Флот", "Ж/д перевозки", "Элеваторы"]


def text(slide, x, y, w, h, s, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, color
    return box


def page(prs, title):
    """Слайд в стиле демо-брендбука: заголовок, полоса-акцент, подвал с периодом."""
    s = prs.slides.add_slide(prs.slide_layouts[6])
    text(s, 0.6, 0.35, 12, 0.8, title, 28, True, NAVY)
    bar = s.shapes.add_shape(1, Inches(0.6), Inches(1.15), Inches(1.2), Inches(0.08))
    bar.fill.solid(); bar.fill.fore_color.rgb = TEAL; bar.line.fill.background()
    text(s, 0.6, 7.0, 9, 0.3, "Управленческий отчёт холдинга · {{период}}", 10, False, GREY)
    return s


def table(slide, name, header, n_rows, x, y, w, h, widths):
    shape = slide.shapes.add_table(n_rows + 1, len(header), Inches(x), Inches(y), Inches(w), Inches(h))
    shape.name = name
    t = shape.table
    for j, wj in enumerate(widths):
        t.columns[j].width = Inches(wj)
    for j, hdr in enumerate(header):
        c = t.cell(0, j)
        c.fill.solid(); c.fill.fore_color.rgb = NAVY
        c.text = hdr
        f = c.text_frame.paragraphs[0].runs[0].font
        f.size, f.bold, f.color.rgb = Pt(14), True, WHITE
    for i in range(1, n_rows + 1):
        for j in range(len(header)):
            c = t.cell(i, j)
            c.fill.solid(); c.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
            c.text = "—"          # прочерк-заглушка: скрипт заменит его, сохранив оформление
            p = c.text_frame.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
            f = p.runs[0].font
            f.size, f.color.rgb = Pt(14), INK
    return shape


def chart(slide, name, kind, categories, series, colors, x, y, w, h):
    data = CategoryChartData()
    data.categories = categories
    for sname, vals in series:
        data.add_series(sname, vals)
    gf = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w), Inches(h), data)
    gf.name = name
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False
    ch.legend.font.size = Pt(12)
    ch.category_axis.tick_labels.font.size = Pt(12)
    ch.value_axis.tick_labels.font.size = Pt(12)
    for s, c in zip(ch.plots[0].series, colors):
        if kind == XL_CHART_TYPE.LINE_MARKERS:
            s.format.line.color.rgb = c
            s.format.line.width = Pt(2.5)
            s.smooth = False
        else:
            s.format.fill.solid(); s.format.fill.fore_color.rgb = c
    return gf


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1. титульный
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(4.6))
    bg.fill.solid(); bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    text(s, 0.8, 1.6, 11, 1.2, "Управленческий отчёт холдинга", 40, True, WHITE)
    text(s, 0.8, 2.9, 11, 0.8, "{{период}}", 26, False, GOLD)
    text(s, 0.8, 5.2, 11.5, 0.8, "Демо-шаблон курса «ИИ без воды». Все цифры выдуманы.", 14, False, GREY)

    # 2. выручка: план и факт
    s = page(prs, "Выручка: план и факт")
    table(s, "Таблица_выручка", ["Направление", "План, млн руб.", "Факт, млн руб.", "Исполнение, %"],
          4, 0.6, 1.6, 12, 3.2, [4.5, 2.5, 2.5, 2.5])
    text(s, 0.6, 5.3, 12, 1.2, "{{итог_выручки}}", 18, False, INK)

    # 3. график план/факт
    s = page(prs, "Выручка по направлениям, млн руб.")
    chart(s, "График_план_факт", XL_CHART_TYPE.COLUMN_CLUSTERED, DIRECTIONS,
          [("План", (1, 1, 1)), ("Факт", (1, 1, 1))], [GREY, TEAL], 0.6, 1.5, 12, 5.3)

    # 4. динамика
    s = page(prs, "Динамика выручки с начала года, млн руб.")
    chart(s, "График_динамика", XL_CHART_TYPE.LINE_MARKERS, ["Янв", "Фев"],
          [(d, (1, 1)) for d in DIRECTIONS], [NAVY, TEAL, GOLD], 0.6, 1.5, 12, 5.3)

    # 5. EBITDA
    s = page(prs, "EBITDA и рентабельность")
    table(s, "Таблица_EBITDA", ["Направление", "Выручка, млн руб.", "EBITDA, млн руб.", "Рентабельность, %"],
          4, 0.6, 1.6, 12, 3.2, [4.5, 2.5, 2.5, 2.5])

    # 6–8. направления
    for d in DIRECTIONS:
        s = page(prs, d)
        table(s, f"Таблица_{d}", ["Показатель", "Прошлый месяц", "Текущий месяц", "Изменение"],
              3, 0.6, 1.6, 12, 2.6, [5.4, 2.2, 2.2, 2.2])
        text(s, 0.6, 4.6, 12, 0.4, "Комментарий дирекции", 14, True, TEAL)
        text(s, 0.6, 5.0, 12, 1.6, "{{комментарий_" + d + "}}", 18, False, INK)

    # 9. долг
    s = page(prs, "Долговая нагрузка")
    table(s, "Таблица_долг", ["Показатель", "Значение"], 4, 0.6, 1.6, 8, 3.2, [5.5, 2.5])
    text(s, 0.6, 5.2, 12, 0.8, "{{статус_ковенанта}}", 22, True, NAVY)

    # 10. выводы
    s = page(prs, "Выводы")
    text(s, 0.6, 1.6, 12, 4.5, "{{выводы}}", 20, False, INK)

    # 11. служебный слайд для проверяющего
    s = page(prs, "Контроль сборки (служебный слайд)")
    text(s, 0.6, 1.6, 12, 5.0, "{{контроль}}", 14, False, INK)

    prs.save(HERE / "template.pptx")
    print("Готово: template.pptx")


if __name__ == "__main__":
    main()
