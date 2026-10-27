# -*- coding: utf-8 -*-
"""Ежемесячная презентация: исходники Excel + шаблон pptx → готовая презентация.

Запуск:   python build_presentation.py data/2026-09
Результат: out/Отчёт_2026-09.pptx

Порядок работы (как в спецификации spec.md):
  1. проверить исходники — если структура изменилась, остановиться с понятной ошибкой;
  2. посчитать показатели;
  3. заполнить шаблон: метки {{…}} в тексте, таблицы и графики по именам;
  4. сохранить файл и вывести контрольные цифры.
Скрипт ничего не рисует и не придумывает: оформление берётся из шаблона,
цифры — из исходников, комментарии подразделений вставляются дословно.
"""
import sys
from pathlib import Path

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData

HERE = Path(__file__).parent
TEMPLATE = HERE / "template.pptx"
OUT_DIR = HERE / "out"

DIRECTIONS = ["Флот", "Ж/д перевозки", "Элеваторы"]
SHEET = {"Флот": "Флот", "Ж/д перевозки": "ЖД", "Элеваторы": "Элеваторы"}
DIVISION_ROWS = {
    "Флот": ["Судов в эксплуатации, шт.", "Загрузка флота, %", "Средняя ставка фрахта, тыс. руб./сут"],
    "Ж/д перевозки": ["Вагонов в парке, шт.", "Оборот вагона, сут", "Простой в ремонте, вагоно-сут"],
    "Элеваторы": ["Принято зерна, тыс. т", "Отгружено зерна, тыс. т", "Загрузка мощностей, %"],
}
DEBT_ROWS = ["Долг, млн руб.", "EBITDA за 12 мес., млн руб.", "Ковенант Долг/EBITDA, не более"]
MONTHS = {1: "Январь", 2: "Февраль", 3: "Март", 4: "Апрель", 5: "Май", 6: "Июнь", 7: "Июль",
          8: "Август", 9: "Сентябрь", 10: "Октябрь", 11: "Ноябрь", 12: "Декабрь"}
TOLERANCE = 0.5   # допустимое расхождение итогов из-за округления, млн руб.


# ---------- форматирование чисел по-русски ----------
def fmt(x, digits=0, sign=False):
    """1234.5 → «1 234,5»; sign=True добавляет «+» к положительным."""
    if sign and round(x, digits) == 0:
        return "0"
    s = f"{x:+,.{digits}f}" if sign else f"{x:,.{digits}f}"
    return s.replace(",", " ").replace(".", ",").replace("-", "−")


# ---------- 1. чтение и проверка исходников ----------
class Inputs:
    def __init__(self, folder: Path):
        self.folder = folder
        self.errors = []

    def error(self, msg):
        self.errors.append(msg)

    def sheet(self, wb, name, header, fname):
        """Лист как список строк; проверяет, что лист есть и заголовок не изменился."""
        if name not in wb.sheetnames:
            self.error(f"{fname}: нет листа «{name}» (есть: {', '.join(wb.sheetnames)})")
            return []
        rows = [list(r) for r in wb[name].iter_rows(values_only=True) if any(v is not None for v in r)]
        if not rows or [str(v).strip() if v is not None else "" for v in rows[0][:len(header)]] != header:
            self.error(f"{fname}, лист «{name}»: заголовок изменился. Ожидали: {' | '.join(header)}")
            return []
        return rows[1:]

    def number(self, value, where):
        if isinstance(value, (int, float)):
            return float(value)
        self.error(f"{where}: ожидали число, а там «{value}»")
        return 0.0

    def by_label(self, rows, labels, where):
        """{подпись: строка}; проверяет, что все нужные строки на месте."""
        found = {str(r[0]).strip(): r for r in rows if r and r[0] is not None}
        for lb in labels:
            if lb not in found:
                self.error(f"{where}: нет строки «{lb}»")
        return found

    def load(self):
        f_main, f_comm = self.folder / "upravlenka.xlsx", self.folder / "kommentarii.xlsx"
        for f in (f_main, f_comm):
            if not f.exists():
                self.error(f"Нет файла {f.name} в папке {self.folder}")
        if self.errors:
            return None
        wb = load_workbook(f_main, data_only=True)
        d = {}

        rows = self.sheet(wb, "Сводка", ["Направление", "План выручки, млн руб.", "Факт выручки, млн руб.",
                                         "EBITDA, млн руб."], f_main.name)
        found = self.by_label(rows, DIRECTIONS + ["Итого"], "Лист «Сводка»")
        d["svodka"] = {}
        for name in DIRECTIONS + ["Итого"]:
            if name in found:
                r = found[name]
                d["svodka"][name] = [self.number(v, f"«Сводка», {name}") for v in r[1:4]]

        rows = self.sheet(wb, "Динамика", ["Месяц", *DIRECTIONS], f_main.name)
        d["months"] = [str(r[0]) for r in rows]
        d["dynamics"] = {name: [self.number(r[i + 1], f"«Динамика», {r[0]}") for r in rows]
                         for i, name in enumerate(DIRECTIONS)}

        d["divisions"] = {}
        for name in DIRECTIONS:
            rows = self.sheet(wb, SHEET[name], ["Показатель", "Прошлый месяц", "Текущий месяц"], f_main.name)
            found = self.by_label(rows, DIVISION_ROWS[name], f"Лист «{SHEET[name]}»")
            d["divisions"][name] = [(lb, self.number(found[lb][1], f"«{SHEET[name]}», {lb}"),
                                     self.number(found[lb][2], f"«{SHEET[name]}», {lb}"))
                                    for lb in DIVISION_ROWS[name] if lb in found]

        rows = self.sheet(wb, "Долг", ["Показатель", "Значение"], f_main.name)
        found = self.by_label(rows, DEBT_ROWS, "Лист «Долг»")
        d["debt"] = {lb: self.number(found[lb][1], f"«Долг», {lb}") for lb in DEBT_ROWS if lb in found}

        wb = load_workbook(f_comm, data_only=True)
        rows = self.sheet(wb, "Комментарии", ["Слайд", "Подразделение", "Текст"], f_comm.name)
        found = self.by_label(rows, DIRECTIONS + ["Выводы"], "Файл комментариев")
        d["comments"] = {}
        for key in DIRECTIONS + ["Выводы"]:
            if key in found:
                text = (found[key][2] or "").strip()
                if not text:
                    self.error(f"Файл комментариев: пустой текст для «{key}»")
                d["comments"][key] = (text, found[key][1] or "")

        self.cross_checks(d)
        return d

    def cross_checks(self, d):
        """Сверки между листами: ловят ошибки в самих исходниках."""
        sv = d.get("svodka", {})
        if all(k in sv for k in DIRECTIONS + ["Итого"]):
            for j, col in enumerate(["план выручки", "факт выручки", "EBITDA"]):
                total = sum(sv[n][j] for n in DIRECTIONS)
                if abs(total - sv["Итого"][j]) > TOLERANCE:
                    self.error(f"«Сводка»: {col} — сумма направлений {fmt(total)} ≠ строке «Итого» {fmt(sv['Итого'][j])}")
        dyn = d.get("dynamics", {})
        for name in DIRECTIONS:
            if name in sv and dyn.get(name):
                if abs(dyn[name][-1] - sv[name][1]) > TOLERANCE:
                    self.error(f"Факт выручки «{name}»: в «Сводке» {fmt(sv[name][1])}, "
                               f"а в последнем месяце «Динамики» {fmt(dyn[name][-1])}")


# ---------- 2. расчёт показателей ----------
def compute(d):
    sv = d["svodka"]
    rev_rows, ebitda_rows = [], []
    for name in DIRECTIONS + ["Итого"]:
        plan, fact, ebitda = sv[name]
        rev_rows.append([name, fmt(plan), fmt(fact), fmt(fact / plan * 100, 1)])
        ebitda_rows.append([name, fmt(fact), fmt(ebitda), fmt(ebitda / fact * 100, 1)])
    plan_t, fact_t, _ = sv["Итого"]
    debt, ebitda_ltm, cov = (d["debt"][k] for k in DEBT_ROWS)
    ratio = debt / ebitda_ltm
    status = (f"Ковенант соблюдён: {fmt(ratio, 2)} при пороге {fmt(cov, 1)}, запас {fmt(cov - ratio, 2)}"
              if ratio <= cov else f"Ковенант НАРУШЕН: {fmt(ratio, 2)} при пороге {fmt(cov, 1)}")
    div_rows = {}
    for name, rows in d["divisions"].items():
        div_rows[name] = []
        for lb, prev, cur in rows:
            k = 1 if (prev % 1 or cur % 1) else 0     # знаков после запятой — одинаково во всей строке
            div_rows[name].append([lb, fmt(prev, k), fmt(cur, k), fmt(cur - prev, k, sign=True)])
    return {
        "rev_rows": rev_rows, "ebitda_rows": ebitda_rows, "div_rows": div_rows,
        "debt_rows": [["Долг, млн руб.", fmt(debt)], ["EBITDA за 12 мес., млн руб.", fmt(ebitda_ltm)],
                      ["Долг / EBITDA", fmt(ratio, 2)], ["Ковенант, не более", fmt(cov, 1)]],
        "rev_total": f"Итого по холдингу: факт {fmt(fact_t)} млн руб. при плане {fmt(plan_t)} — "
                     f"исполнение {fmt(fact_t / plan_t * 100, 1)}%.",
        "status": status, "ratio": ratio,
    }


# ---------- 3. заполнение шаблона ----------
def shapes_by_name(prs):
    return {sh.name: sh for slide in prs.slides for sh in slide.shapes}


def replace_tokens(prs, mapping):
    """Заменяет метки {{…}} в текстовых полях. Работает на уровне абзаца: если PowerPoint
    разбил метку на несколько фрагментов, текст собирается в первый фрагмент с его оформлением."""
    left = set()
    for slide in prs.slides:
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                full = "".join(r.text for r in p.runs)
                if "{{" not in full:
                    continue
                new = full
                for key, val in mapping.items():
                    new = new.replace("{{" + key + "}}", val)
                p.runs[0].text = new
                for r in p.runs[1:]:
                    r.text = ""
                if "{{" in new:
                    left.add(new[new.index("{{"):new.index("}}") + 2] if "}}" in new else new)
    return left


def fill_table(shape, rows):
    """Пишет значения в ячейки, заменяя прочерк-заглушку; оформление ячеек сохраняется."""
    t = shape.table
    if len(rows) != len(t.rows) - 1:
        raise ValueError(f"Таблица «{shape.name}»: в шаблоне {len(t.rows) - 1} строк, а данных {len(rows)}")
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            p = t.cell(i, j).text_frame.paragraphs[0]
            if p.runs:
                p.runs[0].text = str(val)
                for r in p.runs[1:]:
                    r.text = ""
            else:
                p.add_run().text = str(val)


def fill_chart(shape, categories, series):
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    shape.chart.replace_data(data)   # оформление серий из шаблона сохраняется


def build(folder: Path):
    inp = Inputs(folder)
    d = inp.load()
    if inp.errors:
        print("Презентация НЕ собрана: исходники не прошли проверку.\n")
        for e in inp.errors:
            print("  •", e)
        print("\nИсправьте исходники (или обновите скрипт, если структура изменилась намеренно) и запустите снова.")
        return 1

    c = compute(d)
    year, month = folder.name.split("-")
    period = f"{MONTHS[int(month)]} {year}"

    prs = Presentation(TEMPLATE)
    shapes = shapes_by_name(prs)
    fill_table(shapes["Таблица_выручка"], c["rev_rows"])
    fill_table(shapes["Таблица_EBITDA"], c["ebitda_rows"])
    fill_table(shapes["Таблица_долг"], c["debt_rows"])
    for name in DIRECTIONS:
        fill_table(shapes[f"Таблица_{name}"], c["div_rows"][name])
    sv = d["svodka"]
    fill_chart(shapes["График_план_факт"], DIRECTIONS,
               [("План", [sv[n][0] for n in DIRECTIONS]), ("Факт", [sv[n][1] for n in DIRECTIONS])])
    fill_chart(shapes["График_динамика"], d["months"], [(n, d["dynamics"][n]) for n in DIRECTIONS])

    control = [
        f"Исходники: {folder.name}/upravlenka.xlsx, {folder.name}/kommentarii.xlsx — проверки пройдены.",
        f"Контрольные суммы: план {fmt(sv['Итого'][0])}, факт {fmt(sv['Итого'][1])}, "
        f"EBITDA {fmt(sv['Итого'][2])} млн руб. — совпадают с суммой направлений.",
        f"Долг / EBITDA = {fmt(c['ratio'], 2)}.",
        "Комментарии подразделений вставлены дословно из файла комментариев.",
        "Перед отправкой: сверить 2–3 цифры со слайдов с исходниками и удалить этот слайд.",
    ]
    mapping = {"период": period, "итог_выручки": c["rev_total"], "статус_ковенанта": c["status"],
               "выводы": d["comments"]["Выводы"][0], "контроль": "\n".join("• " + x for x in control)}
    for name in DIRECTIONS:
        text, author = d["comments"][name]
        mapping["комментарий_" + name] = f"{text} ({author})"
    left = replace_tokens(prs, mapping)
    if left:
        print("Внимание: в шаблоне остались незаполненные метки:", ", ".join(sorted(left)))

    OUT_DIR.mkdir(exist_ok=True)
    out = OUT_DIR / f"Отчёт_{folder.name}.pptx"
    prs.save(out)
    print(f"Готово: {out}")
    for line in control[1:3]:
        print("  ", line)
    return 0


if __name__ == "__main__":
    if len(sys.argv) != 2:
        print("Использование: python build_presentation.py <папка с исходниками, например data/2026-09>")
        sys.exit(2)
    sys.exit(build(Path(sys.argv[1])))
